"""Independent numerical checks against the released supplementary workbook."""
from __future__ import annotations

import hashlib
import json
import re

import numpy as np
from openpyxl import load_workbook

from .analysis import CLASS_NAMES, SAMPLE_NAMES, composite_table, variance_table
from .dataio import load_bundle, resource_path


def validate() -> dict:
    errors = []
    checks = 0

    def check(condition, label):
        nonlocal checks
        checks += 1
        if not condition:
            errors.append(label)

    paper = load_bundle("paper")
    ml = load_bundle("ml")
    names = json.loads(resource_path("compound_names.json").read_text(encoding="utf-8"))
    wb = load_workbook(resource_path("Supplementary Tables.xlsx"), data_only=True)
    provenance = json.loads(resource_path("provenance.json").read_text(encoding="utf-8"))
    for record in provenance["resources"]:
        path = resource_path(record["resource"])
        check(hashlib.sha256(path.read_bytes()).hexdigest() == record["sha256"], f"Resource checksum: {record['resource']}")

    for key, sheet, samples, counts in [
        ("vol_var", "Table S3", ["HN", "BT", "HT"], {"Esters": 65, "Alcohols": 25, "Acids": 12, "Aldehydes and ketones": 12, "Others": 11}),
        ("vol_heat", "Table S8", ["WB", "HK", "ZZ", "SZ"], {"Esters": 61, "Alcohols": 26, "Acids": 13, "Terpenes": 5, "Aldehydes and ketones": 21, "Alkanes": 6, "Others": 28}),
    ]:
        frame = paper[key].set_index("编号")
        actual_counts = {CLASS_NAMES[k]: int(v) for k, v in frame["类别"].value_counts().items()}
        check(actual_counts == counts, f"Chemical-class counts: {key}")
        matched_codes = set()
        for row in wb[sheet]:
            code = row[0].value
            if code not in frame.index:
                continue
            matched_codes.add(code)
            entry = frame.loc[code]
            check(row[3].value == CLASS_NAMES[entry["类别"]], f"{sheet} class: {code}")
            for j, sample in enumerate(samples):
                cell = row[6 + j]
                match = re.fullmatch(r"([\d.]+)\s*±\s*([\d.]+)([a-z]*)", str(cell.value))
                mean, sd, letter = entry[f"{sample}_均值"], entry[f"{sample}_SD"], entry[f"{sample}_字母"]
                if match:
                    ok = np.allclose([float(match[1]), float(match[2])], [mean, sd], atol=1e-8, rtol=0) and match[3] == letter
                else:
                    ok = bool(np.isnan(mean))
                check(ok, f"{sheet}!{cell.coordinate}: mean/SD/letter")
                threshold = entry.get("统一阈值", entry["阈值"])
                if np.isfinite(mean) and np.isfinite(threshold) and threshold > 0:
                    check(np.isclose(mean / threshold, entry[f"{sample}_OAV"], rtol=0, atol=1e-8), f"{key} {code} {sample}: OAV")
        check(len(matched_codes) == len(frame), f"Complete compound coverage: {sheet}")

        oav_sheet = "Table S5" if key == "vol_var" else "Table S9"
        mapping = {names.get(row["化合物"], row["化合物"]): row for _, row in frame.iterrows()}
        for row in wb[oav_sheet]:
            name = str(row[0].value).rstrip("†")
            if name not in mapping:
                continue
            entry = mapping[name]
            threshold = row[1].value
            for j, sample in enumerate(samples):
                mean = entry[f"{sample}_均值"]
                if not np.isfinite(mean):
                    continue
                estimated_oav = mean / threshold
                roav = 100 * estimated_oav / frame[f"{sample}_OAV"].max()
                for cell, expected in [(row[2 + j], estimated_oav), (row[2 + len(samples) + j], roav)]:
                    check(isinstance(cell.value, (int, float)) and abs(cell.value - expected) <= .0000501, f"{oav_sheet}!{cell.coordinate}: OAV/ROAV")

    variance = variance_table()
    for i, row in variance.iterrows():
        for j, column in enumerate(variance.columns[1:], start=2):
            cell = wb["Table S2"].cell(i + 4, j)
            check(isinstance(cell.value, (int, float)) and abs(cell.value - row[column]) <= .0000501, f"Table S2!{cell.coordinate}: variance fraction")
    check(abs(variance.iloc[-1]["eta2_method"] - .8503012093679885) < 1e-10, "Mean method eta-squared")

    scores = composite_table()
    check(np.allclose(scores.score, scores.stored_score, atol=1e-10, rtol=0), "All fixed-weight score arithmetic")
    score_rows = {row[0].value: row[0].row for row in wb["Table S6"]}
    mode_labels = {
        "全浆发酵": "Whole-mash fermentation",
        "浸渍后发酵": "Post-maceration fermentation",
        "清汁发酵": "Clear-juice fermentation",
    }
    for key in ["z_var", "z_heat", "z_mode"]:
        frame = paper[key]
        zcols = [c for c in frame if c.startswith("Z_")]
        for _, row in frame.iterrows():
            sample = mode_labels.get(row["样品"], str(row["样品"]))
            for old, new in SAMPLE_NAMES.items():
                sample = sample.replace(old, new)
            check(sample in score_rows, f"Table S6 sample coverage: {sample}")
            if sample not in score_rows:
                continue
            for j, value in enumerate([*[row[c] for c in zcols], row["综合评分"]], start=2):
                cell = wb["Table S6"].cell(score_rows[sample], j)
                check(isinstance(cell.value, (int, float)) and abs(cell.value - value) <= 1e-10, f"Table S6!{cell.coordinate}: Z-score/composite")

    shap_blocks = [ml["xgb_med"]["shap"][:, :-1]] + [ml["xgb_met"]["shap"][:, k, :-1] for k in range(4)]
    starts = [row[0].row + 1 for row in wb["Table S7"] if row[0].value == "Sample"]
    check(len(starts) == len(shap_blocks) == 5, "SHAP table block count")
    for start, matrix in zip(starts, shap_blocks):
        for i in range(20):
            for j in range(9):
                cell = wb["Table S7"].cell(start + i, j + 2)
                check(isinstance(cell.value, (int, float)) and abs(cell.value - matrix[i, j]) <= .0000501, f"Table S7!{cell.coordinate}: archived SHAP")
    wb.close()
    return {
        "passed": not errors, "checks_run": checks, "errors": errors,
        "scope": "Consistency of released processed data, supplementary tables and saved model arrays; not validation of raw experiments or model retraining.",
    }
